"""Validate feed.json and write the Merchant Center upload file (TSV) + review workbook (XLSX).
Usage: python3 -I finalize.py <feed.json> <outdir>
"""
import csv, json, re, sys, os
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1:3]
d = json.load(open(SRC, encoding='utf-8'))
feed, skipped = d['feed'], d['skipped']

COLS = ['id', 'title', 'description', 'link', 'image_link', 'additional_image_link', 'availability', 'price',
        'sale_price', 'condition', 'brand', 'gtin', 'mpn', 'identifier_exists', 'google_product_category',
        'product_type']

errors = []


def err(row, msg):
    errors.append(f"{row.get('id')}: {msg}")


def gtin_ok(s):
    if not re.fullmatch(r'\d{8}|\d{12,14}', s):
        return False
    d = [int(c) for c in s]
    tot = sum(x * (3 if i % 2 == 0 else 1) for i, x in enumerate(d[:-1][::-1]))
    return (10 - tot % 10) % 10 == d[-1]


ids = Counter(r['id'] for r in feed)
for r in feed:
    for k in COLS:
        v = r.get(k, '')
        if any(ch in v for ch in '\t\r\n'):
            err(r, f'{k} contains tab/newline')
        r[k] = v.lstrip('"').strip()
    r['description'] = r['description'].lstrip('•.;: ')
    for k in ('id', 'title', 'description', 'link', 'image_link', 'availability', 'price', 'condition', 'brand',
              'google_product_category'):
        if not r[k]:
            err(r, f'missing {k}')
    if ids[r['id']] > 1:
        err(r, 'duplicate id')
    if len(r['id']) > 50:
        err(r, 'id > 50')
    if len(r['title']) > 150:
        err(r, 'title > 150')
    if len(r['description']) > 5000:
        err(r, 'description > 5000')
    if len(r['brand']) > 70:
        err(r, 'brand > 70')
    if not r['link'].startswith('https://superi.ge/'):
        err(r, 'bad link')
    if not re.match(r'https://\S+\.(jpe?g|png|webp|gif|bmp|tiff?)$', r['image_link'], re.I):
        err(r, 'bad image_link ' + r['image_link'])
    extra = [u for u in r['additional_image_link'].split(',') if u]
    if len(extra) > 10 or any(not u.startswith('https://') or ',' in u for u in extra):
        err(r, 'bad additional_image_link')
    if not re.fullmatch(r'\d+\.\d{2} GEL', r['price']):
        err(r, 'bad price')
    if r['sale_price']:
        if not re.fullmatch(r'\d+\.\d{2} GEL', r['sale_price']):
            err(r, 'bad sale_price')
        elif float(r['sale_price'].split()[0]) >= float(r['price'].split()[0]):
            err(r, 'sale_price >= price')
    if r['availability'] not in ('in_stock', 'out_of_stock', 'preorder', 'backorder'):
        err(r, 'bad availability')
    if r['gtin'] and not gtin_ok(r['gtin']):
        err(r, 'bad gtin')
    if r['identifier_exists'] not in ('', 'no'):
        err(r, 'bad identifier_exists')
    if r['identifier_exists'] == 'no' and (r['gtin'] or r['mpn']):
        err(r, 'identifier_exists=no but identifiers present')

if errors:
    print('VALIDATION ERRORS', len(errors))
    for e in errors[:50]:
        print('  ', e)
    sys.exit(1)

os.makedirs(OUT, exist_ok=True)
tsv = os.path.join(OUT, 'superi_google_merchant_feed.tsv')
with open(tsv, 'w', encoding='utf-8', newline='') as f:
    f.write('\t'.join(COLS) + '\n')
    for r in feed:
        f.write('\t'.join(r[k] for k in COLS) + '\n')

# ---------------- review workbook
wb = Workbook()
HEAD = Font(name='Arial', bold=True, color='FFFFFF')
FILL = PatternFill('solid', start_color='1F4E78')
BODY = Font(name='Arial', size=10)


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(vertical='center')
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


ws = wb.active
ws.title = 'feed'
sheet(ws, COLS, ([r[k] for k in COLS] for r in feed),
      [14, 60, 60, 40, 40, 30, 12, 14, 14, 10, 18, 16, 18, 12, 12, 50])

ws = wb.create_sheet('ინსტრუქცია')
INFO = [
    ['რა არის ამ ფაილში'],
    ['ფურცელი „feed“ — Merchant Center-ში ასატვირთი მონაცემები (იგივეა, რაც superi_google_merchant_feed.tsv ფაილში).'],
    ['ასატვირთად გამოიყენეთ .tsv ფაილი (Products → Add products → Add products from a file). ენა: ქართული, ქვეყანა: საქართველო, ვალუტა: GEL.'],
    ['დანარჩენი ფურცლები მხოლოდ შესამოწმებლადაა და Merchant Center-ში არ იტვირთება.'],
    [''],
    ['რა შეიცვალა'],
    ['• დარჩა მხოლოდ აქტიური (A) პროდუქტები; H და D სტატუსის პროდუქტები წაიშალა.'],
    ['• სვეტებს დაერქვა Google-ის ოფიციალური ატრიბუტების სახელები (id, title, description, link, image_link, price ...).'],
    ['• google_product_category — თითოეულ პროდუქტს მიენიჭა Google-ის ოფიციალური კატეგორიის ID (ტაქსონომია 2021-09-21).'],
    ['• product_type — საიტის კატეგორია „ ა > ბ > გ “ ფორმატში.'],
    ['• brand — ახალი სვეტი; ბრენდი აღებულია საიტიდან და შემოწმებულია დასახელებასთან. სადაც ბრენდი არ აქვს, ჩაწერილია მაღაზიის სახელი „Superi“ და identifier_exists = no (Google-ის წესის მიხედვით).'],
    ['• price / sale_price — თუ საიტზე ძველი (გადახაზული) ფასია, price = ძველი ფასი, sale_price = მიმდინარე ფასი.'],
    ['• availability — აღებულია საიტის გვერდიდან (in_stock / out_of_stock).'],
    ['• description — HTML გასუფთავებულია, დარჩა მხოლოდ ტექსტი (მაქს. 5000 სიმბოლო).'],
    ['• image_link — სრული ზომის სურათი საიტიდან; additional_image_link — დამატებითი სურათები (მაქს. 10).'],
    ['• gtin — მხოლოდ შემოწმებული (საკონტროლო ციფრით) შტრიხკოდები; mpn — მწარმოებლის მოდელის კოდი.'],
    ['• title — გასწორდა ბეჭდვითი შეცდომები, დასახელებას წინ დაემატა პროდუქტის ტიპი, ბრენდი ჩაიწერა სწორი ფორმით.'],
    [''],
    ['Merchant Center-ში დამატებით გასაკეთებელი'],
    ['• მიწოდების (Shipping) და დაბრუნების (Returns) პარამეტრები უნდა დააყენოთ Merchant Center-ის ანგარიშში — ფაილში ეს ველები არ არის.'],
]
for row in INFO:
    ws.append(row)
ws.column_dimensions['A'].width = 160
for row in ws.iter_rows():
    for c in row:
        c.font = Font(name='Arial', size=11, bold=c.value in ('რა არის ამ ფაილში', 'რა შეიცვალა', 'Merchant Center-ში დამატებით გასაკეთებელი'))
        c.alignment = Alignment(wrap_text=True, vertical='top')

ws = wb.create_sheet('კატეგორიების რუკა')
cm = Counter((r['product_type'], r['google_product_category'], r['_gpc_path']) for r in feed)
sheet(ws, ['საიტის კატეგორია (product_type)', 'Google ID', 'Google კატეგორია', 'პროდუქტების რაოდენობა'],
      ([a, int(b), c, n] for (a, b, c), n in sorted(cm.items())), [70, 12, 90, 14])

ws = wb.create_sheet('ბრენდები')
bc = Counter(r['brand'] for r in feed)
src = {}
for r in feed:
    src.setdefault(r['brand'], Counter())[r['_brand_src']] += 1
names = {'site': 'საიტის ბრენდი', 'title': 'სათაურიდან ამოღებული', 'title-fix': 'სათაურით შესწორებული',
         'title-line': 'სათაურიდან (ქვე-ბრენდი)', 'store': 'ბრენდის გარეშე → მაღაზიის სახელი'}
sheet(ws, ['ბრენდი', 'პროდუქტების რაოდენობა', 'წყარო'],
      ([b, n, ', '.join(f'{names[k]}: {v}' for k, v in src[b].items())] for b, n in bc.most_common()), [28, 14, 60])

ws = wb.create_sheet('სათაურის ცვლილებები')
sheet(ws, ['id', 'ძველი დასახელება', 'ახალი title'],
      ([r['id'], r['_orig_name'], r['title']] for r in feed if r['_orig_name'] != r['title']), [14, 80, 80])

ws = wb.create_sheet('ამოღებული პროდუქტები')
sheet(ws, ['Product code', 'Product name', 'Product URL', 'მიზეზი'],
      ([s['Product code'], s['Product name'], s['Product URL'], s['მიზეზი']] for s in skipped), [16, 70, 50, 40])

xlsx = os.path.join(OUT, 'superi_google_merchant_feed.xlsx')
wb.save(xlsx)

st = Counter()
for r in feed:
    st['total'] += 1
    st['avail_' + r['availability']] += 1
    st['with_gtin'] += bool(r['gtin'])
    st['with_mpn'] += bool(r['mpn'])
    st['ident_no'] += r['identifier_exists'] == 'no'
    st['on_sale'] += bool(r['sale_price'])
    st['store_brand'] += r['_brand_src'] == 'store'
    st['with_additional_images'] += bool(r['additional_image_link'])
print(json.dumps(st, ensure_ascii=False))
print('skipped', Counter(s['მიზეზი'] for s in skipped))
print('wrote', tsv, xlsx)
